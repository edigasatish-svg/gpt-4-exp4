from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# Create document
doc = Document()

# Add Title
title = doc.add_heading('Ex. No.:4 PERSONALIZED RECOMMENDATION SYSTEM', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_placeholder = doc.add_paragraph('<DATE>                                                                                           <REG. NO>')

# Aim
doc.add_heading('Aim:', level=1)
doc.add_paragraph("To develop a personalized recommendation system using GPT-4 that analyzes user preferences and generates, ranks and explains suitable recommendations based on the user's profile, interests, skill level and goals.")

# Problem Statement
doc.add_heading('Problem Statement:', level=1)
doc.add_paragraph("Design and develop an intelligent personalized recommendation system using GPT-4. The application should collect user information and preferences, analyze the user's requirements, generate personalized recommendations, rank them according to suitability, provide explanations for the recommendations and refine the results based on user feedback.")

# Software Requirements
doc.add_heading('Software Requirements:', level=1)
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Software'
hdr_cells[1].text = 'Version'
reqs = [
    ('Python', '3.10+'),
    ('Ollama', 'Latest'),
    ('VS Code', 'Latest'),
    ('Pandas', '2.0+'),
    ('Sentence Transformers', '2.2+')
]
for item, version in reqs:
    row_cells = table.add_row().cells
    row_cells[0].text = item
    row_cells[1].text = version

# Input Parameters
doc.add_heading('Input Parameters:', level=1)
table2 = doc.add_table(rows=1, cols=2)
table2.style = 'Table Grid'
hdr_cells = table2.rows[0].cells
hdr_cells[0].text = 'Parameters'
hdr_cells[1].text = 'Description'
params = [
    ('Name', "User's full name (String)"),
    ('Age', "User's age (Integer)"),
    ('Background', "Educational or professional background (String)"),
    ('Interests', "Specific topics of interest like AI, Python (String)"),
    ('Skill Level', "Current proficiency - Beginner/Intermediate/Advanced (String)"),
    ('Preferred Category', "Type of learning - Practical, Theoretical (String)"),
    ('Goal', "Ultimate objective to achieve (String)"),
    ('Preferences', "Additional requirements (String)"),
    ('Number of items', "Count of recommendations requested (Integer)")
]
for param, desc in params:
    row_cells = table2.add_row().cells
    row_cells[0].text = param
    row_cells[1].text = desc

# System Architecture
doc.add_heading('System Architecture:', level=1)
doc.add_paragraph("<Students should insert the architecture diagram.>\n[ Architecture Flow: User Profile Input -> Preference Analysis -> Prompt Generation -> GPT-4 / Ollama LLM -> Recommendation Ranking -> Explanation -> User Feedback -> Refined Recommendations. ]")

# Module Design
doc.add_heading('Module Design:', level=1)
table3 = doc.add_table(rows=1, cols=2)
table3.style = 'Table Grid'
hdr_cells = table3.rows[0].cells
hdr_cells[0].text = 'Module'
hdr_cells[1].text = 'Responsibility'
modules = [
    ('Input Module', 'Collects user profile data and preferences.'),
    ('Profile Module', 'Structures and stores user profile attributes.'),
    ('Preference Module', 'Analyzes specific user interests and preferred categories.'),
    ('Prompt Module', 'Dynamically generates a structured prompt for the LLM.'),
    ('Recommendation Module', 'Interacts with the LLM (GPT-4/Ollama) to fetch recommendations.'),
    ('Ranking Module', 'Computes relevance scores and ranks the generated items.'),
    ('Explanation Module', 'Extracts and formats reasoning provided by the LLM.'),
    ('Feedback Module', 'Captures user feedback on initial recommendations.'),
    ('Intent Module', 'Interprets feedback to adjust future recommendations.'),
    ('Save Module', 'Exports the final recommendations to a file (CSV).'),
    ('Main Module', 'Orchestrates the flow between all modules.')
]
for mod, resp in modules:
    row_cells = table3.add_row().cells
    row_cells[0].text = mod
    row_cells[1].text = resp

# Implementation
doc.add_heading('Implementation:', level=1)
doc.add_paragraph("<GitHub Link>:\n<LinkedIn Post Link>:")

doc.add_heading('Source Code (recommendation_system.py):', level=2)
try:
    with open('c:\\gpt 4 lab works\\recommendation_system.py', 'r') as f:
        code = f.read()
    doc.add_paragraph(code)
except FileNotFoundError:
    doc.add_paragraph("[Code snippet missing]")

# Output Screenshots
doc.add_heading('Output Screenshots:', level=1)
doc.add_paragraph("Include screenshots for:")
bullets = [
    "Home Screen", "User Profile Input", "Collected User Profile", 
    "Preference Analysis", "Prompt Generation", "GPT-4 Recommendation", 
    "Recommendation Ranking", "Recommendation Explanation", 
    "User Feedback", "Refined Recommendations", "Intent Detection", 
    "Recommendation Saved Successfully"
]
for b in bullets:
    doc.add_paragraph(b, style='List Bullet')

doc.add_paragraph("\n[ Output Screenshots Placeholder: Insert actual console/UI screenshots here ]")

# Result
doc.add_heading('Result:', level=1)
doc.add_paragraph("Thus, a personalized recommendation system using GPT-4 was successfully developed. The system collects user information, analyzes preferences, generates personalized recommendations, ranks and explains the recommendations, accepts user feedback and produces refined recommendations.")

# Save Document
doc.save("c:\\gpt 4 lab works\\Lab_Record_Ex4_Recommendation_System.docx")
print("Document generated successfully at Lab_Record_Ex4_Recommendation_System.docx")
